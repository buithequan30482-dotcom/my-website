import io
import zipfile
from datetime import datetime, time, timedelta

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from core.models import AuditLog
from heritage.models import DiTich, KhuVuc, LoaiHinh
from interactions.models import DanhGia, GopY, LuotXem, TrangThaiDuyet, YeuThich

from . import services as sv

TZ = timezone.get_current_timezone()


def lt(d, h=10):
    """datetime có múi giờ địa phương."""
    return datetime.combine(d, time(h), tzinfo=TZ)


class DuLieuMau(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.today = timezone.localdate()
        cls.tu, cls.den = cls.today - timedelta(days=9), cls.today          # kỳ 10 ngày
        cls.kv1 = KhuVuc.objects.create(ten="Cái Rồng", slug="cai-rong")
        cls.kv2 = KhuVuc.objects.create(ten="Quan Lạn", slug="quan-lan", la_dao=True)
        lh = LoaiHinh.objects.create(ten="Lịch sử", slug="lich-su")
        cls.dt1 = DiTich.objects.create(ma="VD01", ten="Đình Quan Lạn", cap_xep_hang="qg", khu_vuc=cls.kv2,
                                        dia_chi_moi="Quan Lạn", tom_tat="Đình làng.")
        cls.dt2 = DiTich.objects.create(ma="VD02", ten="Đền Cặp Tiên", cap_xep_hang="tinh", khu_vuc=cls.kv1,
                                        dia_chi_moi="Cái Rồng", tom_tat="Đền thờ.")
        cls.dt1.loai_hinh.add(lh)
        cls.dt2.loai_hinh.add(lh)
        d = cls.today - timedelta(days=2)
        p = cls.tu - timedelta(days=3)  # thuộc kỳ trước
        A, P, H, S = (TrangThaiDuyet.APPROVED, TrangThaiDuyet.PENDING, TrangThaiDuyet.HIDDEN, TrangThaiDuyet.SPAM)
        mk = lambda dt, sao, tt, ngay, **k: DanhGia.objects.create(  # noqa: E731
            di_tich=dt, so_sao=sao, noi_dung="Nhận xét", ten_hien_thi="Khách", trang_thai=tt, ngay_tao=lt(ngay), **k)
        # Kỳ này: 5 đánh giá không spam (4 duyệt: 5,4,3,5) + 1 spam
        mk(cls.dt1, 5, A, d, email="nguyenvan@gmail.com", dien_thoai="0912345678")
        mk(cls.dt1, 4, A, d)
        mk(cls.dt2, 3, A, d)
        mk(cls.dt2, 5, A, d, ngon_ngu="en")
        cls.cho = mk(cls.dt2, 2, P, d, nghi_spam=True, ly_do_nghi="Có liên kết")
        mk(cls.dt1, 1, S, d)
        # Kỳ trước: 2 đánh giá duyệt
        mk(cls.dt1, 4, A, p)
        mk(cls.dt1, 2, A, p)
        # Góp ý kỳ này: 4 (1 xong sau 2 ngày, 1 đóng sau 4 ngày, 1 mới, 1 xuống cấp đang xử lý)
        GopY.objects.create(loai="dich_vu", di_tich=cls.dt1, noi_dung="a", trang_thai="xong", ngay_tao=lt(d - timedelta(days=2)),
                            ngay_xu_ly=lt(d))
        GopY.objects.create(loai="hoi_dap", noi_dung="b", trang_thai="dong", ngay_tao=lt(d - timedelta(days=4)),
                            ngay_xu_ly=lt(d))
        GopY.objects.create(loai="mong_muon", noi_dung="c", trang_thai="moi", ngay_tao=lt(d))
        GopY.objects.create(loai="xuong_cap", di_tich=cls.dt2, noi_dung="d", trang_thai="dang", ngay_tao=lt(d))
        GopY.objects.create(loai="khac", noi_dung="cũ", trang_thai="moi", ngay_tao=lt(p))
        # Lượt xem: kỳ này 100 + 50 + 30; kỳ trước 90
        for ngay, dt, nn, tb, so in [(d, cls.dt1, "vi", "mobile", 100), (d, cls.dt2, "en", "desktop", 50),
                                     (cls.today, None, "vi", "mobile", 30), (p, cls.dt1, "vi", "mobile", 90)]:
            LuotXem.objects.create(ngay=ngay, duong_dan=f"/{nn}/{ngay}/{dt.pk if dt else 0}", di_tich=dt,
                                   ngon_ngu=nn, thiet_bi=tb, so_luot=so)
        YeuThich.objects.create(di_tich=cls.dt1, ma_trinh_duyet="b1", ngay_tao=lt(d))
        YeuThich.objects.create(di_tich=cls.dt1, ma_trinh_duyet="b2", ngay_tao=lt(d))
        YeuThich.objects.create(di_tich=cls.dt2, ma_trinh_duyet="b3", ngay_tao=lt(p))

        User = get_user_model()
        cls.admin = User.objects.create_superuser("quantri", "qt@example.com", "MatKhauRatManh!2025")
        cls.staff = User.objects.create_user("bientap", "bt@example.com", "MatKhauRatManh!2025", is_staff=True)

    def bo_loc(self, **k):
        return sv.parse_bo_loc({"ky": "tuy_chon", "tu": self.tu.isoformat(), "den": self.den.isoformat(), **k})

    def login(self, user=None):
        self.client.force_login(user or self.admin)
        s = self.client.session
        s["otp_ok"] = True
        s.save()


class KpiTests(DuLieuMau):
    def test_kpi_ky_nay(self):
        k = sv.kpi(self.bo_loc())["nay"]
        self.assertEqual(k["luot_xem"], 180)
        self.assertEqual(k["so_danh_gia"], 5)               # trừ spam
        self.assertEqual(k["diem_tb"], 4.25)                # (5+4+3+5)/4
        self.assertEqual(k["ti_le_4_5"], 75.0)              # 3/4
        self.assertEqual(k["gop_y_moi"], 4)
        self.assertEqual(k["ti_le_xu_ly"], 50.0)            # xong + đóng = 2/4
        self.assertEqual(k["tg_xu_ly"], 3.0)                # (2 + 4) / 2 ngày
        self.assertEqual(k["muon_den"], 2)

    def test_so_sanh_ky_truoc(self):
        k = {x["key"]: x for x in sv.kpi(self.bo_loc())["ds"]}
        self.assertEqual(k["luot_xem"]["ky_truoc"], 90)
        self.assertEqual(k["luot_xem"]["thay_doi"], 100.0)
        self.assertEqual(k["diem_tb"]["ky_truoc"], 3.0)
        self.assertTrue(k["luot_xem"]["tot"])

    def test_bo_loc(self):
        self.assertEqual(sv.kpi(self.bo_loc(khu_vuc=str(self.kv1.pk)))["nay"]["so_danh_gia"], 3)
        self.assertEqual(sv.kpi(self.bo_loc(ngon_ngu="en"))["nay"]["luot_xem"], 50)
        self.assertEqual(sv.kpi(self.bo_loc(cap="qg"))["nay"]["luot_xem"], 100)
        self.assertEqual(sv.kpi(self.bo_loc(di_tich=str(self.dt2.pk)))["nay"]["muon_den"], 0)

    def test_gom_nhom_thoi_gian(self):
        self.assertEqual(sv.do_chi_tiet(self.bo_loc()), "ngay")
        self.assertEqual(len(sv.chuoi_thoi_gian(self.bo_loc())["moc"]), 10)
        dai = sv.parse_bo_loc({"ky": "tuy_chon", "tu": "2024-01-01", "den": "2024-06-30"})
        self.assertEqual(sv.do_chi_tiet(dai), "tuan")
        rat_dai = sv.parse_bo_loc({"ky": "tuy_chon", "tu": "2023-01-01", "den": "2024-06-30"})
        self.assertEqual(sv.chuoi_thoi_gian(rat_dai)["muc"], "thang")
        self.assertEqual(len(sv.chuoi_thoi_gian(rat_dai)["moc"]), 18)

    def test_che_thong_tin(self):
        self.assertEqual(sv.che_email("nguyenvan@gmail.com"), "ng***@gmail.com")
        self.assertEqual(sv.che_sdt("0912345678"), "0912***678")

    def test_viec_can_lam(self):
        self.assertEqual(sv.viec_can_lam(), {"cho_duyet": 1, "gop_y_moi": 2, "xuong_cap": 1})


class ViewTests(DuLieuMau):
    def test_index_superuser(self):
        self.login()
        r = self.client.get(reverse("dashboard:index"), {"ky": "30"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Thống kê tương tác")
        self.assertContains(r, 'id="vd-data"')

    def test_index_bo_loc_tuy_chon(self):
        self.login()
        r = self.client.get(reverse("dashboard:index"), {"ky": "tuy_chon", "tu": "2020-01-01", "den": "2026-12-31"})
        self.assertEqual(r.status_code, 200)

    def test_staff_khong_quyen(self):
        self.login(self.staff)
        self.assertEqual(self.client.get(reverse("dashboard:index")).status_code, 403)
        self.staff.user_permissions.add(Permission.objects.get(codename="view_danhgia"))
        self.assertEqual(self.client.get(reverse("dashboard:index")).status_code, 200)

    def test_chua_dang_nhap(self):
        self.assertEqual(self.client.get(reverse("dashboard:index")).status_code, 302)

    def test_export_xlsx(self):
        self.login()
        r = self.client.get(reverse("dashboard:export_xlsx"),
                            {"ky": "tuy_chon", "tu": self.tu.isoformat(), "den": self.den.isoformat()})
        self.assertEqual(r.status_code, 200)
        self.assertIn(f"BaoCao_TuongTac_DiTichVanDon_{self.tu:%Y%m%d}_{self.den:%Y%m%d}.xlsx", r["Content-Disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        for ten in ["Tổng quan", "Theo di tích", "Đánh giá chi tiết", "Góp ý chi tiết", "Lượt xem theo ngày"]:
            self.assertIn(ten, wb.sheetnames)
        ws = wb["Đánh giá chi tiết"]
        header = [c.value for c in ws[1]]
        self.assertNotIn("ip_hash", " ".join(map(str, header)))
        self.assertEqual(ws.max_row, 1 + 6)
        emails = [c.value for c in ws["H"]]
        self.assertIn("ng***@gmail.com", emails)
        self.assertNotIn("nguyenvan@gmail.com", emails)
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        self.assertIn("xl/charts/chart1.xml", charts)
        self.assertGreaterEqual(len(charts), 6)
        self.assertTrue(AuditLog.objects.filter(action="export", user=self.admin).exists())

    def test_export_csv(self):
        self.login()
        for bang in ("danhgia", "gopy", "luotxem"):
            r = self.client.get(reverse("dashboard:export_csv"), {"bang": bang, "ky": "30"})
            self.assertEqual(r.status_code, 200)
            body = b"".join(r.streaming_content).decode("utf-8")
            self.assertTrue(body.startswith("﻿"))
            self.assertIn({"danhgia": "Số sao", "gopy": "Mã tra cứu", "luotxem": "Số lượt"}[bang], body)
            self.assertNotIn("nguyenvan@gmail.com", body)


class ModerationTests(DuLieuMau):
    def test_trang_hang_doi(self):
        self.login()
        r = self.client.get(reverse("dashboard:moderation"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Có liên kết")

    def test_duyet(self):
        self.login()
        r = self.client.post(reverse("dashboard:moderation_action"), {"act": f"approve:{self.cho.pk}"})
        self.assertEqual(r.status_code, 302)
        self.cho.refresh_from_db()
        self.assertEqual(self.cho.trang_thai, TrangThaiDuyet.APPROVED)
        self.assertEqual(self.cho.nguoi_duyet, self.admin)
        self.assertIsNotNone(self.cho.ngay_duyet)
        self.assertTrue(AuditLog.objects.filter(action="approve", object_id=str(self.cho.pk)).exists())

    def test_hang_loat_fetch_va_tra_loi(self):
        self.login()
        r = self.client.post(reverse("dashboard:moderation_action"),
                             {"act": "reply:%d" % self.cho.pk, f"phan_hoi_{self.cho.pk}": "Cảm ơn bạn."},
                             headers={"X-Requested-With": "fetch"})
        self.assertEqual(r.json()["ids"], [self.cho.pk])
        self.cho.refresh_from_db()
        self.assertEqual(self.cho.phan_hoi, "Cảm ơn bạn.")
        moi = DanhGia.objects.create(di_tich=self.dt1, so_sao=3, noi_dung="x", ten_hien_thi="y")
        r = self.client.post(reverse("dashboard:moderation_action"), {"act": "spam:selected", "ids": [moi.pk]},
                             headers={"X-Requested-With": "fetch"})
        self.assertTrue(r.json()["ok"])
        moi.refresh_from_db()
        self.assertEqual(moi.trang_thai, TrangThaiDuyet.SPAM)

    def test_khong_quyen_sua(self):
        self.staff.user_permissions.add(Permission.objects.get(codename="view_danhgia"))
        self.login(self.staff)
        r = self.client.post(reverse("dashboard:moderation_action"), {"act": f"approve:{self.cho.pk}"})
        self.assertEqual(r.status_code, 403)


@override_settings(DEBUG=True)
class SeedTests(DuLieuMau):
    def test_seed_va_xoa(self):
        call_command("seed_demo_interactions", "--danh-gia", "40", "--gop-y", "8", "--ngay", "10", stdout=io.StringIO())
        self.assertEqual(DanhGia.objects.filter(ten_hien_thi__startswith="[DEMO]").count(), 40)
        self.assertEqual(GopY.objects.filter(ho_ten__startswith="[DEMO]").count(), 8)
        self.assertTrue(LuotXem.objects.filter(duong_dan__startswith="/__demo__/").exists())
        call_command("seed_demo_interactions", "--clear", stdout=io.StringIO())
        self.assertFalse(DanhGia.objects.filter(ten_hien_thi__startswith="[DEMO]").exists())
        self.assertEqual(DanhGia.objects.count(), 8)  # dữ liệu thật giữ nguyên
