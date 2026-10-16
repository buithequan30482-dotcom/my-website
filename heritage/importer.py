"""Nhập / xuất hàng loạt di tích bằng Excel (TR FR-A08; bài hướng dẫn "nhap-excel").

Quy trình: build_template() -> người dùng điền -> parse_workbook() trả bảng xem trước (lỗi từng dòng, không ghi CSDL)
-> người dùng chọn "Thêm mới" / "Cập nhật theo mã" -> apply_rows() ghi các dòng hợp lệ trong một transaction.
Di tích thêm mới luôn ở trạng thái Nháp. export_workbook() xuất danh sách theo đúng định dạng mẫu (nhập lại được).
"""
import io
import re
from decimal import Decimal

from django.db import transaction
from django.utils import translation
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.utils import CoordinateError, in_van_don, parse_coordinates, strip_accents

from .models import CapXepHang, DiemDiTich, DiTich, DoTinCay, KhuVuc, LoaiHinh, TrangThai

SHEET = "Di tich"
GUIDE_SHEET = "Huong dan"
LIST_SHEET = "Danh muc"
EXAMPLE_CODE = "VIDU01"
MAX_ROWS = 2000

# (khóa, tiêu đề cột, bắt buộc, độ rộng, ghi chú hướng dẫn)
COLUMNS = [
    ("ma", "Mã", True, 10, "Mã di tích, viết liền, không trùng. Ví dụ VD45."),
    ("ten", "Tên (VI)", True, 34, "Tên chính thức bằng tiếng Việt."),
    ("ten_en", "Tên (EN)", False, 28, "Tên tiếng Anh (không bắt buộc)."),
    ("ten_zh", "Tên (中文)", False, 20, "Tên tiếng Trung giản thể (không bắt buộc)."),
    ("ten_khac", "Tên gọi khác", False, 22, "Tên dân gian, tên chữ… cách nhau dấu phẩy (không bắt buộc)."),
    ("cap", "Cấp xếp hạng", True, 18, "Chọn trong danh sách: Quốc gia đặc biệt / Quốc gia / Cấp tỉnh / Kiểm kê."),
    ("loai_hinh", "Loại hình", True, 28, "Một hoặc nhiều loại, cách nhau dấu phẩy. Xem danh mục ở sheet Huong dan."),
    ("khu_vuc", "Khu vực", True, 16, "Chọn trong danh sách khu vực / đảo."),
    ("dia_chi_cu", "Địa điểm trước 01/7/2025", False, 30, "Địa chỉ cũ (xã, huyện)."),
    ("dia_chi_moi", "Địa điểm từ 01/7/2025", True, 34, "Địa chỉ theo đơn vị hành chính mới."),
    ("tom_tat", "Tóm tắt (VI)", True, 50, "1–3 câu, tối đa 600 ký tự."),
    ("so_quyet_dinh", "Số quyết định", False, 24, "Ví dụ: 313/VH-QĐ ngày 28/4/1962."),
    ("nam", "Năm xếp hạng/kiểm kê", False, 12, "Năm 4 chữ số, ví dụ 1962."),
    ("dien_tich", "Diện tích", False, 14, "Ghi kèm đơn vị, ví dụ 1.200 m²."),
    ("toa_do", "Tọa độ", False, 50,
     "Một hoặc nhiều điểm theo mẫu  Tên điểm: tọa độ; Tên điểm 2: tọa độ. Tọa độ dạng 20°53'49.3\"N 107°18'33.7\"E "
     "hoặc 20.897028, 107.309361. Chỉ một điểm thì có thể bỏ phần tên."),
    ("ghi_chu", "Ghi chú", False, 30, "Ghi chú cho người nhập, hệ thống không lưu cột này."),
]
REQUIRED = [c[0] for c in COLUMNS if c[2]]
LABELS = {c[0]: c[1] for c in COLUMNS}

CAP_LABELS = {CapXepHang.QGDB: "Quốc gia đặc biệt", CapXepHang.QG: "Quốc gia", CapXepHang.TINH: "Cấp tỉnh",
              CapXepHang.KIEMKE: "Kiểm kê"}

EXAMPLE = {
    "ma": EXAMPLE_CODE, "ten": "Đình Ví Dụ (dòng ví dụ – xóa trước khi nhập)", "ten_en": "Vi Du Communal House",
    "ten_zh": "示例亭", "ten_khac": "Đình làng Ví Dụ", "cap": "Cấp tỉnh", "loai_hinh": "Lịch sử, Kiến trúc nghệ thuật",
    "khu_vuc": "", "dia_chi_cu": "Xã Quan Lạn, huyện Vân Đồn",
    "dia_chi_moi": "Thôn Quan Lạn, đặc khu Vân Đồn, tỉnh Quảng Ninh",
    "tom_tat": "Ngôi đình cổ thờ các vị có công giữ đất, giữ làng. Đây là dòng ví dụ – hãy xóa trước khi nhập.",
    "so_quyet_dinh": "1234/QĐ-UBND ngày 01/01/2010", "nam": 2010, "dien_tich": "1.200 m²",
    "toa_do": "Cổng đình: 20°53'49.3\"N 107°18'33.7\"E; Hậu cung: 20.897100, 107.309500",
    "ghi_chu": "Dòng ví dụ",
}

HEAD_FILL = PatternFill("solid", fgColor="0E6B6B")
REQ_FILL = PatternFill("solid", fgColor="A3302A")
EX_FILL = PatternFill("solid", fgColor="F6F0E4")


def _norm(s):
    """Chuẩn hóa để so khớp: bỏ dấu, chữ thường, giữ chữ số / chữ Hán."""
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9㐀-鿿 ]+", " ", strip_accents(str(s or "")))).strip()


HEADER_MAP = {_norm(c[1]): c[0] for c in COLUMNS}
HEADER_MAP.update({_norm(k): v for k, v in {
    "Mã di tích": "ma", "Tên": "ten", "Tên di tích": "ten", "Tên EN": "ten_en", "Tên ZH": "ten_zh",
    "Tên (ZH)": "ten_zh", "Tên khác": "ten_khac", "Cấp": "cap", "Khu vực / đảo": "khu_vuc",
    "Tóm tắt": "tom_tat", "Năm": "nam", "Năm xếp hạng": "nam", "Tọa độ (nhiều điểm)": "toa_do",
}.items()})

CAP_ALIASES = {}
for _code, _label in CAP_LABELS.items():
    for _a in (_label, "di tich " + _label, _code):
        CAP_ALIASES[_norm(_a)] = str(_code)
CAP_ALIASES.update({"qgdb": "qgdb", "dac biet": "qgdb", "tinh": "tinh", "da kiem ke": "kiemke",
                    "kiem ke": "kiemke", "quoc gia dac biet": "qgdb"})


# ---------------------------------------------------------------- File mẫu / xuất
def _catalogs():
    with translation.override("vi"):
        khu_vuc = [k.ten_vi or k.ten for k in KhuVuc.objects.order_by("thu_tu", "ten")]
        loai_hinh = [l.ten_vi or l.ten for l in LoaiHinh.objects.order_by("thu_tu")]
    return khu_vuc, loai_hinh


def _new_workbook():
    khu_vuc, loai_hinh = _catalogs()
    caps = list(CAP_LABELS.values())
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for i, (key, label, req, width, note) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=label + ("*" if req else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQ_FILL if req else HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(note + (" (Bắt buộc)" if req else ""), "He thong")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"

    # Sheet danh mục (ẩn) cho danh sách thả xuống
    dm = wb.create_sheet(LIST_SHEET)
    for col, (title, values) in enumerate((("Cấp xếp hạng", caps), ("Khu vực", khu_vuc),
                                           ("Loại hình", loai_hinh)), start=1):
        dm.cell(row=1, column=col, value=title)
        for r, v in enumerate(values, start=2):
            dm.cell(row=r, column=col, value=v)
    dm.sheet_state = "hidden"
    col = {c[0]: get_column_letter(i) for i, c in enumerate(COLUMNS, start=1)}
    last = MAX_ROWS + 1
    dv_cap = DataValidation(type="list", formula1=f"='{LIST_SHEET}'!$A$2:$A${len(caps) + 1}", allow_blank=True,
                            showErrorMessage=True, errorTitle="Cấp xếp hạng",
                            error="Chọn một giá trị trong danh sách: " + " / ".join(caps))
    dv_cap.add(f"{col['cap']}2:{col['cap']}{last}")
    ws.add_data_validation(dv_cap)
    if khu_vuc:
        dv_kv = DataValidation(type="list", formula1=f"='{LIST_SHEET}'!$B$2:$B${len(khu_vuc) + 1}", allow_blank=True,
                               showErrorMessage=True, errorTitle="Khu vực", error="Chọn khu vực trong danh sách.")
        dv_kv.add(f"{col['khu_vuc']}2:{col['khu_vuc']}{last}")
        ws.add_data_validation(dv_kv)
    dv_lh = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showInputMessage=True,
                           promptTitle="Loại hình", prompt="Cách nhau dấu phẩy: " + ", ".join(loai_hinh)[:240])
    dv_lh.add(f"{col['loai_hinh']}2:{col['loai_hinh']}{last}")
    ws.add_data_validation(dv_lh)
    dv_nam = DataValidation(type="whole", operator="between", formula1="1000", formula2="2100", allow_blank=True,
                            showErrorMessage=True, errorTitle="Năm", error="Nhập năm 4 chữ số, ví dụ 1962.")
    dv_nam.add(f"{col['nam']}2:{col['nam']}{last}")
    ws.add_data_validation(dv_nam)
    dv_tt = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="600", allow_blank=True,
                           showErrorMessage=True, errorTitle="Tóm tắt", error="Tóm tắt tối đa 600 ký tự.")
    dv_tt.add(f"{col['tom_tat']}2:{col['tom_tat']}{last}")
    ws.add_data_validation(dv_tt)

    # Sheet hướng dẫn
    hd = wb.create_sheet(GUIDE_SHEET)
    hd.column_dimensions["A"].width = 30
    hd.column_dimensions["B"].width = 100
    rows = [
        ("HƯỚNG DẪN NHẬP DI TÍCH TỪ EXCEL", ""),
        ("", ""),
        ("Bước 1", f"Điền mỗi di tích một dòng ở sheet “{SHEET}”, từ dòng 2. Cột tiêu đề màu đỏ (dấu *) là bắt buộc."),
        ("Bước 2", f"Dòng 2 là ví dụ (mã {EXAMPLE_CODE}) – xóa trước khi nhập. Nếu quên, hệ thống tự bỏ qua dòng này."),
        ("Bước 3", "Lưu file (.xlsx) → Quản trị → Di tích → Nhập từ Excel → chọn file → Kiểm tra file."),
        ("Bước 4", "Xem bảng kiểm tra: dòng đỏ là lỗi (không được nhập). Sửa trong Excel rồi tải lại nếu cần."),
        ("Bước 5", "Chọn “Thêm mới” hoặc “Cập nhật theo mã” → bấm Xác nhận nhập. Dữ liệu chỉ ghi sau bước này."),
        ("", ""),
        ("Thêm mới", "Chỉ thêm các mã chưa có trong hệ thống; dòng có mã đã tồn tại báo lỗi và không được nhập."),
        ("Cập nhật theo mã", "Mã đã có: cập nhật các ô có dữ liệu (ô trống giữ nguyên giá trị cũ). "
                             "Tọa độ: cập nhật điểm cùng tên, thêm điểm mới. Mã chưa có: thêm mới."),
        ("Trạng thái", "Di tích thêm mới luôn ở trạng thái Nháp. Cần thêm ảnh, chọn ảnh đại diện rồi gửi duyệt."),
        ("Xuất – sửa – nhập lại", "File “Xuất Excel” ở danh sách di tích có cùng định dạng: có thể sửa rồi nhập lại "
                                  "bằng chế độ “Cập nhật theo mã”."),
        ("", ""),
        ("CÁC CỘT", ""),
    ]
    rows += [(label + (" *" if req else ""), note) for _k, label, req, _w, note in COLUMNS]
    rows += [("", ""), ("DANH MỤC", ""), ("Cấp xếp hạng", " / ".join(caps)),
             ("Khu vực", ", ".join(khu_vuc) or "(chưa có danh mục)"),
             ("Loại hình", ", ".join(loai_hinh) or "(chưa có danh mục)"),
             ("", ""), ("VÍ DỤ TỌA ĐỘ", ""),
             ("Một điểm, độ-phút-giây", "20°53'49.3\"N 107°18'33.7\"E"),
             ("Một điểm, thập phân", "20.897028, 107.309361"),
             ("Nhiều điểm", "Đình: 20°53'49.3\"N 107°18'33.7\"E; Chùa: 20.8981, 107.3102")]
    for r, (a, b) in enumerate(rows, start=1):
        hd.cell(row=r, column=1, value=a).font = Font(bold=True, color="0A4A4E")
        cb = hd.cell(row=r, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
    hd["A1"].font = Font(bold=True, size=14, color="0E6B6B")
    wb.active = 0
    return wb, ws, khu_vuc, loai_hinh


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template():
    """Sinh file mẫu .xlsx (bytes) với dòng 2 là ví dụ."""
    wb, ws, khu_vuc, loai_hinh = _new_workbook()
    example = dict(EXAMPLE)
    if khu_vuc:
        example["khu_vuc"] = khu_vuc[-1] if "Quan Lạn" not in khu_vuc else "Quan Lạn"
    if loai_hinh:
        example["loai_hinh"] = ", ".join(loai_hinh[:2])
    for i, (key, *_r) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=2, column=i, value=example.get(key, ""))
        c.font = Font(italic=True, color="6F675B")
        c.fill = EX_FILL
    ws.cell(row=2, column=1).comment = Comment(f"Dòng ví dụ – XÓA DÒNG NÀY trước khi nhập. Nếu quên, hệ thống tự bỏ "
                                               f"qua dòng có mã {EXAMPLE_CODE}.", "He thong")
    return _save(wb)


def format_points(points):
    """Danh sách DiemDiTich -> 'Tên: tọa độ; Tên: tọa độ' (dạng nhập lại được)."""
    parts = []
    for p in points:
        if p.vi_do is None or p.kinh_do is None:
            continue
        coord = p.toa_do_goc or f"{p.vi_do}, {p.kinh_do}"
        try:
            parse_coordinates(coord)
        except CoordinateError:
            coord = f"{p.vi_do}, {p.kinh_do}"
        name = (p.ten_vi or p.ten or "").replace(";", ",").replace(":", " –").strip()
        parts.append(f"{name}: {coord}" if name else coord)
    return "; ".join(parts)


def export_rows(queryset):
    with translation.override("vi"):
        for d in queryset.select_related("khu_vuc").prefetch_related("loai_hinh", "diem"):
            yield {
                "ma": d.ma, "ten": d.ten_vi or d.ten, "ten_en": d.ten_en or "", "ten_zh": d.ten_zh_hans or "",
                "ten_khac": d.ten_khac, "cap": CAP_LABELS.get(d.cap_xep_hang, d.cap_xep_hang),
                "loai_hinh": ", ".join(l.ten_vi or l.ten for l in d.loai_hinh.all()),
                "khu_vuc": d.khu_vuc.ten_vi or d.khu_vuc.ten, "dia_chi_cu": d.dia_chi_cu,
                "dia_chi_moi": d.dia_chi_moi_vi or d.dia_chi_moi, "tom_tat": d.tom_tat_vi or d.tom_tat,
                "so_quyet_dinh": d.so_quyet_dinh, "nam": d.nam_xep_hang, "dien_tich": d.dien_tich,
                "toa_do": format_points(d.diem.all()), "ghi_chu": f"Trạng thái: {d.get_trang_thai_display()}",
            }


def export_workbook(queryset):
    """Xuất danh sách di tích theo đúng định dạng file mẫu (bytes)."""
    wb, ws, *_ = _new_workbook()
    for r, row in enumerate(export_rows(queryset), start=2):
        for i, (key, *_x) in enumerate(COLUMNS, start=1):
            v = row.get(key)
            ws.cell(row=r, column=i, value=v if v not in (None, "") else None)
    return _save(wb)


# ---------------------------------------------------------------- Đọc file
def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if hasattr(v, "year") and hasattr(v, "month"):  # ô ngày tháng
        return str(v.year)
    return str(v).strip()


def parse_points(text, default_name):
    """'Tên: tọa độ; Tên: tọa độ' -> ([{ten, lat, lng, goc, note}], [lỗi], [cảnh báo])."""
    pts, errors, warns = [], [], []
    if not text:
        return pts, errors, warns
    for part in (p.strip() for p in str(text).replace("\n", ";").split(";")):
        if not part:
            continue
        name, coord = "", part
        if ":" in part:
            name, coord = (x.strip() for x in part.rsplit(":", 1))
        if not name:
            name = default_name if not pts else f"{default_name} – điểm {len(pts) + 1}"
        try:
            lat, lng, note = parse_coordinates(coord)
        except CoordinateError:
            errors.append(f"Tọa độ “{part[:60]}” sai định dạng (đúng: 20°53'49.3\"N 107°18'33.7\"E hoặc "
                          "20.897028, 107.309361).")
            continue
        if not in_van_don(lat, lng):
            warns.append(f"Điểm “{name}” ({lat}, {lng}) nằm ngoài phạm vi Vân Đồn – kiểm tra lại.")
        pts.append({"ten": name[:300], "lat": lat, "lng": lng, "goc": coord[:120], "note": note})
    return pts, errors, warns


def parse_workbook(fileobj, mode="new"):
    """Đọc file, kiểm tra từng dòng. Trả về {rows: [...], errors: [lỗi chung]}. Không ghi CSDL.

    Mỗi dòng: {row, ma, raw, errors, missing, warnings, skip, exists, ...}. Lỗi phụ thuộc chế độ nhập – "mã đã tồn tại"
    (khi Thêm mới) và "thiếu trường bắt buộc" (không áp dụng khi Cập nhật mã đã có) – được suy ra bởi row_errors().
    """
    try:
        wb = load_workbook(fileobj, data_only=True, read_only=True)
    except Exception:
        return {"rows": [], "errors": ["Không mở được file. Hãy dùng file .xlsx lưu từ Excel theo đúng mẫu."]}
    names = {_norm(n): n for n in wb.sheetnames}
    ws = wb[names[_norm(SHEET)]] if _norm(SHEET) in names else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        return {"rows": [], "errors": ["File trống."]}
    colmap = {}
    for idx, h in enumerate(header or ()):
        key = HEADER_MAP.get(_norm(str(h or "").replace("*", "")))
        if key and key not in colmap:
            colmap[key] = idx
    missing_cols = [LABELS[k] for k in REQUIRED if k not in colmap]
    if missing_cols:
        return {"rows": [], "errors": ["File thiếu cột: " + ", ".join(missing_cols) + ". Hãy tải lại file mẫu."]}

    with translation.override("vi"):
        kv_lookup, lh_lookup = {}, {}
        for k in KhuVuc.objects.all():
            for n in (k.ten_vi, k.ten, k.slug):
                if n:
                    kv_lookup[_norm(n)] = k
        for lo in LoaiHinh.objects.all():
            for n in (lo.ten_vi, lo.ten, lo.slug):
                if n:
                    lh_lookup[_norm(n)] = lo
    existing = {m.upper(): (m, i) for m, i in DiTich.objects.values_list("ma", "id")}

    rows, seen = [], {}
    for rnum, values in enumerate(it, start=2):
        if rnum > MAX_ROWS + 1:
            wb.close()
            return {"rows": [], "errors": [f"File có quá {MAX_ROWS} dòng. Hãy chia nhỏ file."]}
        values = list(values or ())
        raw = {k: _cell_str(values[i]) if i < len(values) else "" for k, i in colmap.items()}
        if not any(v for k, v in raw.items() if k != "ghi_chu"):
            continue
        r = {"row": rnum, "raw": raw, "errors": [], "missing": [], "warnings": [], "skip": False, "exists": False,
             "ma": raw.get("ma", "")}
        if raw.get("ma", "").upper() == EXAMPLE_CODE:
            r["skip"] = True
            r["warnings"].append("Dòng ví dụ – được bỏ qua.")
            rows.append(r)
            continue
        for k in REQUIRED:
            if not raw.get(k):
                # Khi cập nhật theo mã, ô trống = giữ nguyên giá trị cũ (chỉ bắt buộc với di tích mới).
                (r["errors"] if k == "ma" else r["missing"]).append(f"Thiếu “{LABELS[k]}”.")
        ma = raw.get("ma", "")
        if ma:
            key = ma.upper()
            if len(ma) > 20:
                r["errors"].append("Mã dài quá 20 ký tự.")
            if re.search(r"\s", ma):
                r["errors"].append("Mã không được có dấu cách.")
            if key in seen:
                r["errors"].append(f"Mã {ma} trùng với dòng {seen[key]} trong file.")
            else:
                seen[key] = rnum
            if key in existing:
                r["ma"], r["existing_id"] = existing[key]
                r["exists"] = True
        cap = raw.get("cap", "")
        if cap:
            code = CAP_ALIASES.get(_norm(cap))
            if not code:
                r["errors"].append(f"Cấp xếp hạng “{cap}” không có trong danh mục "
                                   "(Quốc gia đặc biệt / Quốc gia / Cấp tỉnh / Kiểm kê).")
            r["cap"] = code
        kv = raw.get("khu_vuc", "")
        if kv:
            obj = kv_lookup.get(_norm(kv))
            if not obj:
                r["errors"].append(f"Khu vực “{kv}” không có trong danh mục.")
            else:
                r["khu_vuc_id"], r["khu_vuc_ten"] = obj.pk, obj.ten_vi or obj.ten
        lh = raw.get("loai_hinh", "")
        if lh:
            ids, names_ = [], []
            for part in re.split(r"[,;\n]+", lh):
                part = part.strip()
                if not part:
                    continue
                obj = lh_lookup.get(_norm(part))
                if not obj:
                    r["errors"].append(f"Loại hình “{part}” không có trong danh mục.")
                elif obj.pk not in ids:
                    ids.append(obj.pk)
                    names_.append(obj.ten_vi or obj.ten)
            r["loai_hinh_ids"], r["loai_hinh_ten"] = ids, names_
        nam = raw.get("nam", "")
        r["nam"] = None
        if nam:
            if re.fullmatch(r"\d{4}", nam) and 1000 <= int(nam) <= 2100:
                r["nam"] = int(nam)
            else:
                r["errors"].append(f"Năm “{nam}” không hợp lệ (cần 4 chữ số).")
        if len(raw.get("tom_tat", "")) > 600:
            r["errors"].append("Tóm tắt dài quá 600 ký tự.")
        for k, mx in (("ten", 300), ("ten_en", 300), ("ten_zh", 300), ("ten_khac", 300), ("dia_chi_cu", 300),
                      ("dia_chi_moi", 300), ("so_quyet_dinh", 200), ("dien_tich", 300)):
            if len(raw.get(k, "")) > mx:
                r["errors"].append(f"“{LABELS[k]}” dài quá {mx} ký tự.")
        pts, perr, pwarn = parse_points(raw.get("toa_do", ""), raw.get("ten", "") or ma)
        r["points"] = pts
        r["errors"] += perr
        r["warnings"] += pwarn
        if not raw.get("toa_do"):
            r["warnings"].append("Chưa có tọa độ – bổ sung sau ở thẻ Vị trí.")
        rows.append(r)
    wb.close()
    if not rows:
        return {"rows": [], "errors": ["Không có dòng dữ liệu nào (từ dòng 2)."]}
    return {"rows": rows, "errors": []}


EXISTS_ERROR = "Mã {ma} đã có trong hệ thống – chọn “Cập nhật theo mã” nếu muốn cập nhật."


def row_errors(r, mode):
    errs = list(r["errors"])
    if mode == "new" or not r.get("exists"):
        errs = list(r.get("missing", [])) + errs
    if mode == "new" and r.get("exists"):
        errs.append(EXISTS_ERROR.format(ma=r["ma"]))
    return errs


def is_valid(r, mode):
    return not r["skip"] and not row_errors(r, mode)


def summarize(rows, mode="new"):
    data = [r for r in rows if not r["skip"]]
    valid = [r for r in data if is_valid(r, mode)]
    return {
        "total": len(data), "valid": len(valid), "invalid": len(data) - len(valid),
        "new": len([r for r in valid if not r.get("exists")]),
        "existing": len([r for r in valid if r.get("exists")]),
        "warnings": len([r for r in data if r["warnings"]]),
    }


# ---------------------------------------------------------------- Ghi dữ liệu
def _set(obj, name, value, update):
    """Gán trường; ở chế độ cập nhật, ô trống giữ nguyên giá trị cũ."""
    if update and (value is None or value == ""):
        return
    setattr(obj, name, value if value is not None else "")


def apply_rows(rows, mode, user=None):
    """Ghi các dòng hợp lệ. mode: 'new' | 'update'. Trả về {created, updated, skipped} (danh sách mã)."""
    if mode not in ("new", "update"):
        raise ValueError("mode")
    with translation.override("vi"), transaction.atomic():
        return _apply_rows(rows, mode, user)


def _apply_rows(rows, mode, user):
    result = {"created": [], "updated": [], "skipped": []}
    for r in rows:
        if r["skip"] or r["errors"] or (r.get("missing") and not (mode == "update" and r.get("exists"))):
            continue
        raw = r["raw"]
        obj = DiTich.objects.filter(ma__iexact=r["ma"]).first()
        if obj and mode != "update":
            result["skipped"].append(obj.ma)
            continue
        update = obj is not None
        if not obj:
            obj = DiTich(ma=r["ma"], trang_thai=TrangThai.DRAFT, nguoi_tao=user)
        obj.nguoi_sua = user
        for name, value in (("ten_vi", raw.get("ten")), ("ten_en", raw.get("ten_en")),
                            ("ten_zh_hans", raw.get("ten_zh")), ("ten_khac", raw.get("ten_khac")),
                            ("cap_xep_hang", r.get("cap")), ("dia_chi_cu", raw.get("dia_chi_cu")),
                            ("dia_chi_moi_vi", raw.get("dia_chi_moi")), ("tom_tat_vi", raw.get("tom_tat")),
                            ("so_quyet_dinh", raw.get("so_quyet_dinh")), ("nam_xep_hang", r.get("nam")),
                            ("dien_tich", raw.get("dien_tich"))):
            if name == "nam_xep_hang":
                if value is not None or not update:
                    obj.nam_xep_hang = value
                continue
            _set(obj, name, value, update)
        if not update:
            obj.ten = obj.ten_vi
            obj.tom_tat = obj.tom_tat_vi
            obj.dia_chi_moi = obj.dia_chi_moi_vi
        if r.get("khu_vuc_id"):
            obj.khu_vuc_id = r["khu_vuc_id"]
        obj.save()
        if r.get("loai_hinh_ids"):
            obj.loai_hinh.set(r["loai_hinh_ids"])
        existing_pts = {_norm(p.ten_vi or p.ten): p for p in obj.diem.all()} if update else {}
        order = obj.diem.count()
        has_main = obj.diem.filter(diem_chinh=True).exists()
        for p in r.get("points", []):
            d = existing_pts.get(_norm(p["ten"]))
            if not d:
                d = DiemDiTich(di_tich=obj, thu_tu=order, do_tin_cay=DoTinCay.DANH_MUC)
                order += 1
                d.ten_vi = d.ten = p["ten"]
            d.vi_do, d.kinh_do = Decimal(f"{p['lat']:.6f}"), Decimal(f"{p['lng']:.6f}")
            d.toa_do_goc = p["goc"]
            if not has_main:
                d.diem_chinh = has_main = True
            d.save()
        (result["updated"] if update else result["created"]).append(obj.ma)
    return result
