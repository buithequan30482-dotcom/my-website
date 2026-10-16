"""Kiểm thử khu quản trị nội dung di tích: tọa độ, nhập/xuất Excel, tải ảnh, quyền xuất bản, các trang quản trị."""
import io
import shutil
import tempfile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from core.models import AuditLog
from core.utils import CoordinateError, parse_coordinates

from . import importer
from .admin_forms import DiemDiTichInlineForm, DiTichAdminForm
from .models import CapXepHang, DiemDiTich, DiTich, HinhAnh, KhuVuc, LoaiHinh, TrangThai

MEDIA = tempfile.mkdtemp(prefix="vd-test-media-")


def png_bytes(w=800, h=600, color=(14, 107, 107), fmt="PNG"):
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, format=fmt)
    return buf.getvalue()


@override_settings(MEDIA_ROOT=MEDIA)
class BaseAdminTest(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    @classmethod
    def setUpTestData(cls):
        cls.kv = KhuVuc.objects.create(ten="Quan Lạn", ten_vi="Quan Lạn", slug="quan-lan", dao="Quan Lạn", la_dao=True)
        cls.kv2 = KhuVuc.objects.create(ten="Cái Rồng", ten_vi="Cái Rồng", slug="cai-rong")
        cls.lh1 = LoaiHinh.objects.create(ten="Lịch sử", ten_vi="Lịch sử", slug="lich-su", thu_tu=1)
        cls.lh2 = LoaiHinh.objects.create(ten="Kiến trúc nghệ thuật", ten_vi="Kiến trúc nghệ thuật", slug="kien-truc", thu_tu=2)
        cls.d = DiTich.objects.create(ma="VD90", ten_vi="Đình Thử", ten="Đình Thử", cap_xep_hang=CapXepHang.TINH,
                                      khu_vuc=cls.kv, dia_chi_moi_vi="Thôn Quan Lạn", dia_chi_moi="Thôn Quan Lạn",
                                      tom_tat_vi="Tóm tắt thử.", tom_tat="Tóm tắt thử.", trang_thai=TrangThai.DRAFT)
        cls.d.loai_hinh.add(cls.lh1)
        DiemDiTich.objects.create(di_tich=cls.d, ten="Cổng", ten_vi="Cổng", vi_do="20.897028", kinh_do="107.309361")
        U = get_user_model()
        cls.admin = U.objects.create_superuser("sa", "sa@example.com", "x-Pass-123")
        cls.editor = U.objects.create_user("btv", "btv@example.com", "x-Pass-123", is_staff=True)
        perms = Permission.objects.filter(content_type__app_label="heritage", codename__in=[
            "add_ditich", "change_ditich", "view_ditich", "add_hinhanh", "change_hinhanh", "view_hinhanh",
            "add_diemditich", "change_diemditich", "view_diemditich", "add_video", "change_video", "view_video",
            "add_nguonthamkhao", "change_nguonthamkhao", "view_nguonthamkhao"])
        cls.editor.user_permissions.set(perms)

    def login(self, user):
        self.client.force_login(user)
        s = self.client.session
        s["otp_ok"] = True
        s.save()


# ------------------------------------------------------------------ Tọa độ
class CoordinateTests(TestCase):
    def test_dms(self):
        lat, lng, _ = parse_coordinates('20°53\'49.3"N 107°18\'33.7"E')
        self.assertAlmostEqual(lat, 20.897028, places=5)
        self.assertAlmostEqual(lng, 107.309361, places=5)

    def test_dms_seconds_60_normalized(self):
        lat, lng, note = parse_coordinates('20°48\'60.0"N 107°30\'00"E')
        self.assertAlmostEqual(lat, 20.816667, places=5)
        self.assertTrue(note)

    def test_decimal(self):
        self.assertEqual(parse_coordinates("20.897028, 107.309361")[:2], (20.897028, 107.309361))
        self.assertEqual(parse_coordinates("20,8970 107,3093")[:2], (20.897, 107.3093))

    def test_invalid(self):
        for bad in ("", "abc", "20.89", "200, 500", "20°N"):
            with self.assertRaises(CoordinateError):
                parse_coordinates(bad)

    def form(self, data):
        from django.forms import modelform_factory
        F = modelform_factory(DiemDiTich, form=DiemDiTichInlineForm,
                              fields=["ten_vi", "toa_do", "vi_do", "kinh_do", "do_tin_cay", "thu_tu"])
        return F(data=data)

    def test_inline_form_parses_text(self):
        f = self.form({"ten_vi": "Điểm A", "toa_do": '20°53\'49.3"N 107°18\'33.7"E',
                                       "do_tin_cay": "chua", "thu_tu": "0"})
        self.assertTrue(f.is_valid(), f.errors)
        self.assertAlmostEqual(float(f.cleaned_data["vi_do"]), 20.897028, places=5)
        self.assertEqual(f.instance.toa_do_goc, '20°53\'49.3"N 107°18\'33.7"E')

    def test_inline_form_rejects_bad_text(self):
        f = self.form({"ten_vi": "Điểm A", "toa_do": "không phải tọa độ", "do_tin_cay": "chua"})
        self.assertFalse(f.is_valid())
        self.assertIn("toa_do", f.errors)

    def test_inline_form_needs_both(self):
        f = self.form({"ten_vi": "Điểm A", "vi_do": "20.9", "do_tin_cay": "chua"})
        self.assertFalse(f.is_valid())


# ------------------------------------------------------------------ Nhập / xuất Excel
class ImporterTests(BaseAdminTest):
    def _book(self, rows, header=None):
        wb = load_workbook(io.BytesIO(importer.build_template()))
        ws = wb[importer.SHEET]
        ws.delete_rows(2, 1)
        if header:
            for i, h in enumerate(header, start=1):
                ws.cell(row=1, column=i, value=h)
        for r, row in enumerate(rows, start=2):
            for i, (key, *_x) in enumerate(importer.COLUMNS, start=1):
                ws.cell(row=r, column=i, value=row.get(key))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def row(self, **kw):
        base = {"ma": "VD91", "ten": "Chùa Mới", "cap": "Cấp tỉnh", "loai_hinh": "Lịch sử, Kiến trúc nghệ thuật",
                "khu_vuc": "Quan Lạn", "dia_chi_moi": "Thôn Quan Lạn", "tom_tat": "Chùa cổ.",
                "toa_do": "Chính điện: 20°53'49.3\"N 107°18'33.7\"E; Tháp: 20.8981, 107.3102"}
        base.update(kw)
        return base

    def test_template_roundtrip_example_row_skipped(self):
        data = importer.build_template()
        wb = load_workbook(io.BytesIO(data))
        self.assertIn(importer.SHEET, wb.sheetnames)
        self.assertIn(importer.GUIDE_SHEET, wb.sheetnames)
        headers = [c.value for c in wb[importer.SHEET][1]]
        self.assertEqual(headers[0], "Mã*")
        self.assertIn("Tên (中文)", headers)
        self.assertTrue(wb[importer.SHEET].data_validations.dataValidation)
        parsed = importer.parse_workbook(io.BytesIO(data))
        self.assertEqual(parsed["errors"], [])
        self.assertEqual(len(parsed["rows"]), 1)
        self.assertTrue(parsed["rows"][0]["skip"])

    def test_valid_and_error_rows(self):
        buf = self._book([
            self.row(),
            self.row(ma="VD92", cap="Cấp huyện", khu_vuc="Sao Hỏa", loai_hinh="Lịch sử, Không có", toa_do="abc"),
            self.row(ma="VD91"),  # trùng mã trong file
            self.row(ma="VD93", ten="", tom_tat=""),
            self.row(ma="VD94", toa_do="21.9, 105.8"),  # ngoài Vân Đồn -> cảnh báo
            self.row(ma="VD90"),  # đã tồn tại
        ])
        rows = importer.parse_workbook(buf)["rows"]
        self.assertEqual(rows[0]["errors"], [])
        self.assertEqual(len(rows[0]["points"]), 2)
        self.assertEqual(rows[0]["points"][0]["ten"], "Chính điện")
        errs = " ".join(rows[1]["errors"])
        for frag in ("Cấp xếp hạng", "Khu vực", "Loại hình “Không có”", "Tọa độ"):
            self.assertIn(frag, errs)
        self.assertTrue(any("trùng với dòng 2" in e for e in rows[2]["errors"]))
        self.assertTrue(any("Thiếu “Tên (VI)”" in e for e in importer.row_errors(rows[3], "new")))
        self.assertEqual(rows[4]["errors"], [])
        self.assertTrue(any("ngoài phạm vi" in w for w in rows[4]["warnings"]))
        self.assertTrue(rows[5]["exists"])
        self.assertEqual(rows[5]["errors"], [])
        self.assertTrue(importer.row_errors(rows[5], "new"))
        self.assertFalse(importer.row_errors(rows[5], "update"))
        s = importer.summarize(rows, "new")
        self.assertEqual((s["valid"], s["invalid"]), (2, 4))

    def test_missing_columns(self):
        buf = self._book([], header=["Mã", "Linh tinh"] + [""] * 14)
        res = importer.parse_workbook(buf)
        self.assertTrue(res["errors"])

    def test_apply_new_and_update(self):
        rows = importer.parse_workbook(self._book([self.row(), self.row(ma="VD90", ten="Đình Thử (sửa)", tom_tat="")]))["rows"]
        res = importer.apply_rows(rows, "new")
        self.assertEqual(res["created"], ["VD91"])
        self.assertNotIn("VD90", res["created"] + res["updated"])  # dòng lỗi (thiếu tóm tắt) khi Thêm mới
        d = DiTich.objects.get(ma="VD91")
        self.assertEqual(d.trang_thai, TrangThai.DRAFT)
        self.assertEqual(d.ten_vi, "Chùa Mới")
        self.assertEqual(d.loai_hinh.count(), 2)
        self.assertEqual(d.diem.count(), 2)
        self.assertEqual(d.diem.filter(diem_chinh=True).count(), 1)
        self.assertEqual(DiTich.objects.get(ma="VD90").ten_vi, "Đình Thử")
        res = importer.apply_rows(rows, "update")
        self.assertIn("VD90", res["updated"])
        d90 = DiTich.objects.get(ma="VD90")
        self.assertEqual(d90.ten_vi, "Đình Thử (sửa)")
        self.assertEqual(d90.tom_tat_vi, "Tóm tắt thử.")  # ô trống giữ nguyên
        self.assertEqual(d90.diem.count(), 3)  # Cổng + 2 điểm mới
        self.assertEqual(d90.trang_thai, TrangThai.DRAFT)

    def test_export_reimport(self):
        data = importer.export_workbook(DiTich.objects.filter(ma="VD90"))
        rows = importer.parse_workbook(io.BytesIO(data))["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["errors"], [])
        self.assertTrue(rows[0]["exists"])
        self.assertEqual(rows[0]["points"][0]["ten"], "Cổng")

    def test_import_views_flow(self):
        self.login(self.admin)
        url = reverse("admin:heritage_ditich_import")
        self.assertEqual(self.client.get(url).status_code, 200)
        r = self.client.get(reverse("admin:heritage_ditich_import_template"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        f = SimpleUploadedFile("ds.xlsx", self._book([self.row(), self.row(ma="VD95", cap="Sai")]).getvalue())
        r = self.client.post(url, {"file": f, "mode": "new"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Kiểm tra dữ liệu")
        self.assertFalse(DiTich.objects.filter(ma="VD91").exists())  # chưa ghi khi chưa xác nhận
        token = r.context["token"]
        r = self.client.post(url, {"token": token, "mode": "new"})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(DiTich.objects.filter(ma="VD91").exists())
        self.assertFalse(DiTich.objects.filter(ma="VD95").exists())
        self.assertTrue(AuditLog.objects.filter(action="import").exists())
        r = self.client.post(url, {"token": token, "mode": "new"})  # token đã dùng
        self.assertEqual(r.status_code, 302)

    def test_export_action(self):
        self.login(self.admin)
        r = self.client.post(reverse("admin:heritage_ditich_changelist"),
                             {"action": "xuat_excel", "_selected_action": [self.d.pk]})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])


# ------------------------------------------------------------------ Ảnh
class ImageUploadTests(BaseAdminTest):
    def test_upload_png_creates_variants(self):
        self.login(self.editor)
        url = reverse("admin:heritage_ditich_upload", args=[self.d.pk])
        f = SimpleUploadedFile("anh-thu.png", png_bytes(800, 600), content_type="image/png")
        r = self.client.post(url, {"file": f, "nguon": "Ảnh: Phòng VHXH", "tac_gia": "A"})
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        self.assertTrue(data["ok"])
        h = HinhAnh.objects.get(pk=data["id"])
        self.assertEqual(h.di_tich, self.d)
        self.assertEqual(h.nguon, "Ảnh: Phòng VHXH")
        self.assertIn("480", h.bien_the)
        self.assertTrue(data["thumb"].endswith(".webp"))
        self.assertTrue(any("nhỏ hơn" in w for w in data["warnings"]))  # 800px < 1600
        self.d.refresh_from_db()
        self.assertEqual(self.d.anh_dai_dien_id, h.pk)  # ảnh đầu tiên thành ảnh đại diện
        self.assertTrue(data["avatar"])

    def test_upload_rejects_fake_and_missing_source(self):
        self.login(self.editor)
        url = reverse("admin:heritage_ditich_upload", args=[self.d.pk])
        fake = SimpleUploadedFile("gia.jpg", b"not an image at all", content_type="image/jpeg")
        r = self.client.post(url, {"file": fake, "nguon": "X", "tac_gia": "A"})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.json()["ok"])
        ok = SimpleUploadedFile("a.png", png_bytes(), content_type="image/png")
        r = self.client.post(url, {"file": ok, "nguon": "", "tac_gia": "A"})
        self.assertEqual(r.status_code, 400)
        gif = io.BytesIO()
        Image.new("RGB", (10, 10)).save(gif, format="GIF")
        r = self.client.post(url, {"file": SimpleUploadedFile("a.gif", gif.getvalue()), "nguon": "X"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(HinhAnh.objects.count(), 0)

    def test_avatar_order_update_delete(self):
        self.login(self.admin)
        up = reverse("admin:heritage_ditich_upload", args=[self.d.pk])
        ids = []
        for i in range(3):
            r = self.client.post(up, {"file": SimpleUploadedFile(f"a{i}.png", png_bytes(), content_type="image/png"),
                                      "nguon": "N", "tac_gia": "T"})
            ids.append(r.json()["id"])
        r = self.client.post(reverse("admin:heritage_ditich_image_avatar", args=[self.d.pk, ids[2]]))
        self.assertTrue(r.json()["ok"])
        self.d.refresh_from_db()
        self.assertEqual(self.d.anh_dai_dien_id, ids[2])
        r = self.client.post(reverse("admin:heritage_ditich_image_order", args=[self.d.pk]),
                             {"ids": f"{ids[2]},{ids[0]},{ids[1]}"})
        self.assertTrue(r.json()["ok"])
        self.assertEqual(list(self.d.hinh_anh.order_by("thu_tu").values_list("id", flat=True)), [ids[2], ids[0], ids[1]])
        upd = reverse("admin:heritage_ditich_image_update", args=[self.d.pk, ids[0]])
        self.assertEqual(self.client.post(upd, {"nguon": ""}).status_code, 400)
        r = self.client.post(upd, {"chu_thich_vi": "Cổng đình", "chu_thich_en": "Gate"})
        self.assertTrue(r.json()["ok"])
        h = HinhAnh.objects.get(pk=ids[0])
        self.assertEqual((h.chu_thich_vi, h.chu_thich_en), ("Cổng đình", "Gate"))
        r = self.client.post(reverse("admin:heritage_ditich_image_delete", args=[self.d.pk, ids[2]]))
        self.assertTrue(r.json()["was_avatar"])
        self.d.refresh_from_db()
        self.assertIsNone(self.d.anh_dai_dien_id)

    def test_editor_cannot_delete_image(self):
        h = HinhAnh.objects.create(di_tich=self.d, tep=SimpleUploadedFile("x.png", png_bytes()), nguon="N")
        self.login(self.editor)
        r = self.client.post(reverse("admin:heritage_ditich_image_delete", args=[self.d.pk, h.pk]))
        self.assertEqual(r.status_code, 403)

    def test_picker_json(self):
        HinhAnh.objects.create(di_tich=self.d, tep=SimpleUploadedFile("x.png", png_bytes()), nguon="N", chu_thich_vi="Mái")
        self.login(self.editor)
        r = self.client.get(reverse("admin:heritage_hinhanh_picker") + f"?di_tich={self.d.pk}")
        self.assertEqual(r.json()["items"][0]["caption"], "Mái")


# ------------------------------------------------------------------ Quyền, form, trang
class PermissionAndPageTests(BaseAdminTest):
    def _post_data(self, obj, **over):
        data = {
            "ma": obj.ma, "ten_vi": obj.ten_vi, "ten_khac": "", "cap_xep_hang": obj.cap_xep_hang,
            "loai_hinh": [self.lh1.pk], "khu_vuc": obj.khu_vuc_id, "dia_chi_moi_vi": obj.dia_chi_moi_vi,
            "dia_chi_cu": "", "dien_tich": "", "hien_trang": "tot", "tham_quan_duoc": "on", "tom_tat_vi": obj.tom_tat_vi,
            "noi_dung_vi": "<h3>Lịch sử</h3><p>Nội dung <script>alert(1)</script></p>", "thong_tin_tham_quan_vi": "",
            "gio_mo_cua_vi": "", "phi_tham_quan_vi": "Miễn phí", "lien_he": "", "trang_thai": obj.trang_thai,
            "thu_tu": 100, "slug": obj.slug, "so_quyet_dinh": "", "nam_xep_hang": "", "co_quan_qd": "",
            "anh_dai_dien": obj.anh_dai_dien_id or "",
            "diem-TOTAL_FORMS": 1, "diem-INITIAL_FORMS": 1, "diem-MIN_NUM_FORMS": 0, "diem-MAX_NUM_FORMS": 1000,
            "diem-0-id": obj.diem.first().pk, "diem-0-di_tich": obj.pk, "diem-0-ten_vi": "Cổng",
            "diem-0-toa_do": '20°53\'49.3"N 107°18\'33.7"E', "diem-0-vi_do": "20.897028", "diem-0-kinh_do": "107.309361",
            "diem-0-do_tin_cay": "thuc_dia", "diem-0-thu_tu": 1,
            "videos-TOTAL_FORMS": 0, "videos-INITIAL_FORMS": 0, "videos-MIN_NUM_FORMS": 0, "videos-MAX_NUM_FORMS": 1000,
            "nguon-TOTAL_FORMS": 0, "nguon-INITIAL_FORMS": 0, "nguon-MIN_NUM_FORMS": 0, "nguon-MAX_NUM_FORMS": 1000,
        }
        data.update(over)
        return data

    def test_editor_cannot_publish(self):
        self.login(self.editor)
        h = HinhAnh.objects.create(di_tich=self.d, tep=SimpleUploadedFile("x.png", png_bytes()), nguon="N")
        url = reverse("admin:heritage_ditich_change", args=[self.d.pk])
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        choices = [c[0] for c in r.context["adminform"].form.fields["trang_thai"].choices]
        self.assertEqual(sorted(choices), ["draft", "review"])
        r = self.client.post(url, self._post_data(self.d, trang_thai="published", anh_dai_dien=h.pk))
        self.assertEqual(r.status_code, 200)  # form lỗi
        self.d.refresh_from_db()
        self.assertEqual(self.d.trang_thai, TrangThai.DRAFT)
        # action xuất bản không có trong danh sách thao tác của biên tập viên
        r = self.client.get(reverse("admin:heritage_ditich_changelist"))
        self.assertNotIn("xuat_ban", [a[0] for a in r.context["action_form"].fields["action"].choices])
        r = self.client.post(reverse("admin:heritage_ditich_changelist"),
                             {"action": "xuat_ban", "_selected_action": [self.d.pk]})
        self.d.refresh_from_db()
        self.assertEqual(self.d.trang_thai, TrangThai.DRAFT)

    def test_publish_requires_conditions(self):
        self.login(self.admin)
        url = reverse("admin:heritage_ditich_change", args=[self.d.pk])
        r = self.client.post(url, self._post_data(self.d, trang_thai="published"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("ảnh đại diện", str(r.context["adminform"].form.errors["trang_thai"]))
        h = HinhAnh.objects.create(di_tich=self.d, tep=SimpleUploadedFile("x.png", png_bytes()), nguon="N")
        r = self.client.post(url, self._post_data(self.d, trang_thai="published", anh_dai_dien=h.pk))
        self.assertEqual(r.status_code, 302, getattr(r, "context", None) and r.context["adminform"].form.errors)
        self.d.refresh_from_db()
        self.assertEqual(self.d.trang_thai, TrangThai.PUBLISHED)
        self.assertNotIn("<script", self.d.noi_dung_vi)
        p = self.d.diem.first()
        self.assertEqual(p.toa_do_goc, '20°53\'49.3"N 107°18\'33.7"E')

    def test_publish_action_skips_incomplete(self):
        self.login(self.admin)
        r = self.client.post(reverse("admin:heritage_ditich_changelist"),
                             {"action": "xuat_ban", "_selected_action": [self.d.pk]}, follow=True)
        self.d.refresh_from_db()
        self.assertEqual(self.d.trang_thai, TrangThai.DRAFT)
        self.assertContains(r, "ảnh đại diện")

    def test_avatar_choices_only_own_images(self):
        other = DiTich.objects.create(ma="VD99", ten_vi="Khác", ten="Khác", cap_xep_hang="tinh", khu_vuc=self.kv,
                                      dia_chi_moi_vi="x", tom_tat_vi="x")
        mine = HinhAnh.objects.create(di_tich=self.d, tep=SimpleUploadedFile("m.png", png_bytes()), nguon="N")
        theirs = HinhAnh.objects.create(di_tich=other, tep=SimpleUploadedFile("t.png", png_bytes()), nguon="N")
        form = DiTichAdminForm(instance=self.d)
        qs = form.fields["anh_dai_dien"].queryset
        self.assertIn(mine, qs)
        self.assertNotIn(theirs, qs)

    def test_duplicate_coordinate_warning(self):
        other = DiTich.objects.create(ma="VD98", ten_vi="Khác", ten="Khác", cap_xep_hang="tinh", khu_vuc=self.kv,
                                      dia_chi_moi_vi="x", tom_tat_vi="x")
        DiemDiTich.objects.create(di_tich=other, ten_vi="Gần", ten="Gần", vi_do="20.897100", kinh_do="107.309400")
        self.login(self.admin)
        r = self.client.post(reverse("admin:heritage_ditich_change", args=[self.d.pk]), self._post_data(self.d),
                             follow=True)
        self.assertContains(r, "trùng tọa độ")

    def test_change_form_pages(self):
        self.login(self.admin)
        r = self.client.get(reverse("admin:heritage_ditich_change", args=[self.d.pk]))
        self.assertEqual(r.status_code, 200)
        for frag in ("vdh-tabs", "Bản dịch 中文", "SEO &amp; xuất bản", "vdh-map-config", "Chuẩn ảnh", "Đọc tọa độ",
                     "vendor/quill/quill.js", "Xem trước"):
            self.assertContains(r, frag)
        self.assertIn("server.arcgisonline.com", r["Content-Security-Policy"])
        self.assertEqual(self.client.get(reverse("admin:heritage_ditich_changelist")).status_code, 200)
        self.assertEqual(self.client.get(reverse("admin:heritage_ditich_history", args=[self.d.pk])).status_code, 200)
        for name in ("heritage_hinhanh_changelist", "heritage_tuyenthamquan_add", "heritage_lehoi_add",
                     "heritage_thongtinhuuich_changelist", "heritage_tintuc_add", "heritage_trangtinh_add",
                     "heritage_khuvuc_changelist", "heritage_loaihinh_changelist", "heritage_video_add"):
            self.assertEqual(self.client.get(reverse("admin:" + name)).status_code, 200, name)

    def test_preview_unpublished(self):
        self.login(self.admin)
        url = reverse("admin:heritage_ditich_preview", args=[self.d.pk])
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "XEM TRƯỚC")
        self.assertContains(r, "Đình Thử")
        self.assertEqual(r["X-Robots-Tag"], "noindex, nofollow")
        self.client.logout()
        r = self.client.get(url)
        self.assertEqual(r.status_code, 302)  # chưa đăng nhập -> trang đăng nhập

    def test_wizard(self):
        self.login(self.editor)
        r = self.client.get(reverse("admin:heritage_ditich_wizard"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("wizard=1", r["Location"])
        r = self.client.get(r["Location"])
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Kiểm tra và gửi duyệt")
        self.assertContains(r, "Lưu nháp để tải ảnh")
        # Lưu nháp ở bước 2 -> sang bước 3 (tóm tắt chưa bắt buộc)
        data = {
            "ma": "VD97", "ten_vi": "Miếu Thử", "cap_xep_hang": "kiemke", "loai_hinh": [self.lh1.pk], "khu_vuc": self.kv.pk,
            "dia_chi_moi_vi": "Thôn A", "hien_trang": "tot", "tham_quan_duoc": "on", "trang_thai": "draft",
            "thu_tu": 100, "phi_tham_quan_vi": "Miễn phí", "_vd_draft": "1", "_vd_wizard": "3",
            "diem-TOTAL_FORMS": 1, "diem-INITIAL_FORMS": 0, "diem-MIN_NUM_FORMS": 0, "diem-MAX_NUM_FORMS": 1000,
            "diem-0-ten_vi": "Miếu", "diem-0-toa_do": "20.9, 107.4", "diem-0-do_tin_cay": "chua",
            "videos-TOTAL_FORMS": 0, "videos-INITIAL_FORMS": 0, "videos-MIN_NUM_FORMS": 0, "videos-MAX_NUM_FORMS": 1000,
            "nguon-TOTAL_FORMS": 0, "nguon-INITIAL_FORMS": 0, "nguon-MIN_NUM_FORMS": 0, "nguon-MAX_NUM_FORMS": 1000,
        }
        r = self.client.post(reverse("admin:heritage_ditich_add") + "?wizard=1", data)
        self.assertEqual(r.status_code, 302, r.context and r.context["adminform"].form.errors)
        d = DiTich.objects.get(ma="VD97")
        self.assertIn(f"/{d.pk}/change/?wizard=1#step=3", r["Location"])
        self.assertEqual(d.diem.get().toa_do_goc, "20.9, 107.4")
        self.assertEqual(d.nguoi_tao, self.editor)
        # Gửi duyệt
        data.pop("_vd_draft")
        data.update({"_vd_send_review": "1", "tom_tat_vi": "Miếu nhỏ.", "diem-INITIAL_FORMS": 1,
                     "diem-0-id": d.diem.get().pk, "diem-0-di_tich": d.pk, "slug": d.slug})
        r = self.client.post(reverse("admin:heritage_ditich_change", args=[d.pk]) + "?wizard=1", data)
        self.assertEqual(r.status_code, 302)
        d.refresh_from_db()
        self.assertEqual(d.trang_thai, TrangThai.REVIEW)
